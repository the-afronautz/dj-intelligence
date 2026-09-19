"""Offline tests: no network. Fakes the sources layer with real JEFFERY data
plus deliberate conflicts, and checks reconciliation, flags, cache and export.

    python3 test_catalog.py
"""
from __future__ import annotations

import sys, tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))

import camelot, sources, lookup, export

FAILS = []
def check(name, cond, detail=""):
    print(("  PASS  " if cond else "  FAIL  ") + name + (f"   {detail}" if detail and not cond else ""))
    if not cond: FAILS.append(name)

# --------------------------------------------------------------- camelot
print("\nCamelot conversion (verified JEFFERY values)")
JEFFERY = [
    ("Wyclef Jean", "F minor", 132, "4A"), ("Floyd Mayweather", "Ab major", 134, "4B"),
    ("RiRi", "A♭ major", 140, "4B"), ("Guwop", "C minor", 142, "5A"),
    ("Future Swag", "A major", 154, "11B"), ("Kanye West", "F minor", 108, "4A"),
    ("Harambe", "Ab major", 132, "4B"), ("Webbie", "C# minor", 79, "12A"),
    ("Pick Up the Phone", "G minor", 137, "6A"), ("Swizz Beatz", "E minor", 130, "9A"),
]
for title, key, bpm, expected in JEFFERY:
    check(f"{title}: {key} -> {expected}", camelot.to_camelot(key) == expected,
          f"got {camelot.to_camelot(key)}")
check("enharmonic Db major == C# major", camelot.to_camelot("Db major") == camelot.to_camelot("C# major"))
check("garbage returns None", camelot.to_camelot("not a key") is None)
check("empty returns None", camelot.to_camelot("") is None)

# ----------------------------------------------------------- reconcile
print("\nReconciliation")
r = lookup.reconcile([{"source": "getsongbpm", "bpm": 132, "camelot": "4A", "key": "F minor"}])
check("single source is chosen", r["chosen_source"] == "getsongbpm" and r["bpm"] == 132)
check("no spurious flags", r["flags"] == [], str(r["flags"]))

r = lookup.reconcile([{"source": "getsongbpm", "bpm": 79, "camelot": "12A"}])
check("lone 79 BPM flagged half-time",
      any("half-time" in f for f in r["flags"]), str(r["flags"]))
check("half-time flag shows the x2 value",
      any("158" in f for f in r["flags"]), str(r["flags"]))

r = lookup.reconcile([
    {"source": "getsongbpm", "bpm": 79, "camelot": "12A"},
    {"source": "deezer", "bpm": 158},
])
check("79 vs 158 = octave, not conflict",
      any("bpm-octave" in f for f in r["flags"]) and not any("bpm-conflict" in f for f in r["flags"]),
      str(r["flags"]))
check("priority: getsongbpm beats deezer", r["chosen_source"] == "getsongbpm")

r = lookup.reconcile([
    {"source": "getsongbpm", "bpm": 132, "camelot": "4A"},
    {"source": "deezer", "bpm": 96},
])
check("132 vs 96 = real conflict", any("bpm-conflict" in f for f in r["flags"]), str(r["flags"]))

r = lookup.reconcile([
    {"source": "getsongbpm", "bpm": 132, "camelot": "4A"},
    {"source": "acousticbrainz", "bpm": 132, "camelot": "4B"},
])
check("relative major/minor = adjacent, not conflict",
      any("key-adjacent" in f for f in r["flags"]) and not any("key-conflict" in f for f in r["flags"]),
      str(r["flags"]))

r = lookup.reconcile([
    {"source": "getsongbpm", "bpm": 132, "camelot": "4A"},
    {"source": "acousticbrainz", "bpm": 132, "camelot": "9A"},
])
check("distant key = conflict", any("key-conflict" in f for f in r["flags"]), str(r["flags"]))

r = lookup.reconcile([{"source": "deezer", "bpm": 137}, {"source": "manual", "camelot": "6A", "key": "G minor"}])
check("manual key + deezer bpm combine", r["bpm"] == 137 and r["camelot"] == "6A", str(r))

check("empty readings flag no-data", "no-data" in lookup.reconcile([])["flags"])
check("133.5 vs 134 within tolerance",
      not any("conflict" in f for f in lookup.reconcile(
          [{"source": "getsongbpm", "bpm": 134}, {"source": "deezer", "bpm": 133.5}])["flags"]))

# ------------------------------------------------- end-to-end with fakes
print("\nEnd-to-end (faked network)")
sources.find_release = lambda a, b, prefer_year=None: {
    "mbid": "fake-mbid", "title": "Jeffery", "date": "2016-08-26",
    "status": "Official", "artist": a, "candidates": []}
sources.get_tracklist = lambda mbid: [
    {"position": i + 1, "disc": 1, "title": t[0], "recording_mbid": f"rec-{i}",
     "length_ms": 200000, "credited_artists": "Young Thug"}
    for i, t in enumerate(JEFFERY)]
_by_title = {t[0]: t for t in JEFFERY}
sources.from_getsongbpm = lambda title, artist, api_key=None: (
    None if title == "Pick Up the Phone" else
    {"source": "getsongbpm", "bpm": _by_title[title][2],
     "key": camelot.pretty_key(_by_title[title][1]),
     "camelot": camelot.to_camelot(_by_title[title][1])})
sources.from_acousticbrainz = lambda mbid: None
sources.from_deezer = lambda title, artist: (
    {"source": "deezer", "bpm": 158.0} if title == "Webbie" else None)

with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    conn = lookup.open_cache(td / "t.db")
    release, tracks = lookup.lookup_album("Young Thug", "Jeffery", conn, log=lambda *a: None)
    check("10 tracks returned", len(tracks) == 10, str(len(tracks)))
    check("track order preserved", [t.position for t in tracks] == list(range(1, 11)))
    check("9 resolved, 1 gap", sum(1 for t in tracks if t.chosen_source) == 9)
    webbie = tracks[7]
    check("Webbie flagged as octave (deezer says 158)",
          any("bpm-octave" in f for f in webbie.flags), str(webbie.flags))
    putp = tracks[8]
    check("Pick Up the Phone has no data", putp.chosen_source is None and "no-data" in putp.flags)

    # cache round-trip
    _, again = lookup.lookup_album("Young Thug", "Jeffery", conn, log=lambda *a: None)
    check("cache returns same values", [t.bpm for t in again] == [t.bpm for t in tracks])

    # gap fill
    album_id = release["album_id"]
    ok = lookup.apply_manual(conn, album_id, 9, 137.0, "6A", "G minor", source="tunebat")
    check("apply_manual writes", ok)
    _, filled = lookup.lookup_album("Young Thug", "Jeffery", conn, log=lambda *a: None)
    check("gap now filled from tunebat",
          filled[8].bpm == 137.0 and filled[8].camelot == "6A" and filled[8].chosen_source == "tunebat",
          str(filled[8]))
    check("all 10 resolved after fill", sum(1 for t in filled if t.chosen_source) == 10)

    xlsx = export.write_xlsx(td / "out.xlsx", "Young Thug", "Jeffery", release, filled)
    check("xlsx written", xlsx.exists() and xlsx.stat().st_size > 4000, str(xlsx.stat().st_size))
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    check("two sheets incl. Sources", "Sources" in wb.sheetnames, str(wb.sheetnames))
    ws = wb[wb.sheetnames[0]]
    check("xlsx row 5 is track 1", ws.cell(row=5, column=2).value == "Wyclef Jean" and ws.cell(row=5, column=3).value == 132)
    check("xlsx camelot column correct", ws.cell(row=5, column=4).value == "4A")

    gp, n = export.write_gaps_csv(td / "g.csv", "Young Thug", filled)
    check("gaps csv only lists the flagged track", n == 1, f"n={n}")
    check("gaps csv has a tunebat URL", "tunebat.com/Search" in gp.read_text())

    md = export.write_markdown("Young Thug", "Jeffery", filled)
    check("markdown has 10 rows", md.count("\n| ") >= 10)

print("\n" + ("ALL PASS" if not FAILS else f"{len(FAILS)} FAILURE(S): " + ", ".join(FAILS)))
sys.exit(1 if FAILS else 0)
