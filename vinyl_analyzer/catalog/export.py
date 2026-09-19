"""Excel + CSV output, styled to match the existing track-picks sheets."""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import sources

HEADER_FILL = PatternFill("solid", fgColor="1F2A37")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FLAG_FILL = PatternFill("solid", fgColor="FFF3CD")
MISSING_FILL = PatternFill("solid", fgColor="F8D7DA")

COLUMNS = [
    ("#", 5), ("Track", 46), ("BPM", 8), ("Camelot", 9), ("Key", 14),
    ("Source", 15), ("Flags", 40), ("Credited", 40), ("Recording MBID", 38),
]


def write_xlsx(path: Path, artist: str, album: str, release: dict, tracks: Iterable) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = (album or "Album")[:28]

    ws.append([f"{artist} — {album}"])
    ws["A1"].font = Font(bold=True, size=14)
    sub = f"MusicBrainz release {release.get('mbid','?')}"
    if release.get("date"):
        sub += f" · released {release['date']}"
    ws.append([sub])
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([])

    head_row = ws.max_row + 1
    ws.append([c[0] for c in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=head_row, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width

    for t in tracks:
        ws.append([
            t.position, t.title, t.bpm, t.camelot, t.key,
            t.chosen_source or "", "; ".join(t.flags), t.credited or "",
            t.recording_mbid or "",
        ])
        r = ws.max_row
        if not t.chosen_source:
            for c in range(1, len(COLUMNS) + 1):
                ws.cell(row=r, column=c).fill = MISSING_FILL
        elif t.flags:
            for c in range(1, len(COLUMNS) + 1):
                ws.cell(row=r, column=c).fill = FLAG_FILL
        ws.cell(row=r, column=3).number_format = "0.0"
        for c in (1, 3, 4):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")

    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    # Second sheet: every source's raw answer, so a disputed value is auditable
    raw = wb.create_sheet("Sources")
    raw.append(["#", "Track", "Source", "BPM", "Camelot", "Key"])
    for i, cell in enumerate(raw[1], start=1):
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for w, col in zip((5, 46, 16, 8, 9, 14), "ABCDEF"):
        raw.column_dimensions[col].width = w
    for t in tracks:
        for r in (t.raw or []):
            raw.append([t.position, t.title, r.get("source"), r.get("bpm"),
                        r.get("camelot"), r.get("key")])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_gaps_csv(path: Path, artist: str, tracks: Iterable) -> tuple[Path, int]:
    """Rows for every track that has no data or a flagged conflict, each with a
    Tunebat search URL. Fill bpm/camelot in, then feed it back with --merge-gaps.
    """
    rows = [t for t in tracks if not t.chosen_source or not t.camelot or t.flags]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["position", "title", "current_bpm", "current_camelot",
                    "flags", "tunebat_url", "bpm", "camelot"])
        for t in rows:
            w.writerow([t.position, t.title, t.bpm or "", t.camelot or "",
                        "; ".join(t.flags),
                        sources.tunebat_search_url(t.title, artist), "", ""])
    return path, len(rows)


def write_markdown(artist: str, album: str, tracks: Iterable) -> str:
    lines = [f"**{artist} — {album}**", "",
             "| # | Track | BPM | Camelot | Source | Flags |",
             "|---|-------|-----|---------|--------|-------|"]
    for t in tracks:
        lines.append(
            f"| {t.position} | {t.title} | {f'{t.bpm:g}' if t.bpm else '—'} | "
            f"{t.camelot or '—'} | {t.chosen_source or '—'} | {'; '.join(t.flags) or ''} |"
        )
    return "\n".join(lines)
